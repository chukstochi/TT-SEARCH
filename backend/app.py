import os, io, base64, threading, json, hmac, hashlib, re, uuid, time, requests
from dotenv import load_dotenv
from flask import Flask, request, jsonify, send_file, g
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from itsdangerous import URLSafeTimedSerializer, BadSignature
from docx import Document
import anthropic
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt

load_dotenv()
app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL")
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "dev")
CORS(app, origins=[os.getenv("FRONTEND_URL", "*")])
db = SQLAlchemy(app)
signer = URLSafeTimedSerializer(app.config["SECRET_KEY"])
# Which AI to use: "gemini" (free tier) or "anthropic" (paid). Change AI_PROVIDER in .env.
PROVIDER = os.getenv("AI_PROVIDER", "anthropic").strip().lower()
MODEL = os.getenv("AI_MODEL", "claude-sonnet-5-5")            # used when AI_PROVIDER=anthropic
GEMINI_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash").strip()
GEMINI_FALLBACKS = [m.strip() for m in os.getenv("GEMINI_FALLBACK_MODELS", "").split(",") if m.strip()]  # tried when the main model is busy
GEMINI_TIMEOUT = int(os.getenv("GEMINI_TIMEOUT_SECONDS", "300"))   # how long to wait for one long answer
GEMINI_GAP = float(os.getenv("GEMINI_MIN_GAP_SECONDS", "4"))   # pause between calls to respect free limits
# If your Anthropic key is not tied to a workspace, set ANTHROPIC_WORKSPACE_ID in .env
_ws = os.getenv("ANTHROPIC_WORKSPACE_ID", "").strip()
ai = anthropic.Anthropic(default_headers={"anthropic-workspace-id": _ws} if _ws else None) if PROVIDER == "anthropic" else None
PAY_REQUIRED = os.getenv("PAYMENT_REQUIRED", "false").lower() == "true"
PRICE = int(os.getenv("PRICE_NGN", "5000"))
PSK = os.getenv("PAYSTACK_SECRET", "")

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100))
    email = db.Column(db.String(190), unique=True)
    pw = db.Column(db.String(255))

class Project(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"))
    level = db.Column(db.String(10)); department = db.Column(db.String(120))
    university = db.Column(db.String(160)); title = db.Column(db.String(300))
    topic = db.Column(db.JSON); content = db.Column(db.JSON)  # content: {section: text}
    paid = db.Column(db.Boolean, default=False)
    status = db.Column(db.String(12), default='idle'); progress = db.Column(db.Integer, default=0)
    diagrams = db.Column(db.JSON); error = db.Column(db.String(200))
    created = db.Column(db.DateTime, server_default=db.func.now())

class Saved(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'))
    level = db.Column(db.String(10)); department = db.Column(db.String(120)); university = db.Column(db.String(160))
    topic = db.Column(db.JSON)

class Payment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    reference = db.Column(db.String(64), unique=True)
    project_id = db.Column(db.Integer, db.ForeignKey("project.id"))
    amount = db.Column(db.Integer); status = db.Column(db.String(20), default="pending")

# ---------- helpers ----------
# Shows the real reason (bad key, no credit, ...) instead of a blank 500 when an AI call fails
@app.errorhandler(anthropic.APIStatusError)
def ai_error(e):
    try:
        msg = e.body["error"]["message"]
    except Exception:
        msg = str(e)
    print("[AI error]", msg)
    return jsonify(error=f"AI service error: {msg}"), 502

class AIError(Exception):
    pass

@app.errorhandler(AIError)
def ai_error2(e):
    print("[AI error]", e)
    return jsonify(error=f"AI service error: {e}"), 502

def auth(f):
    from functools import wraps
    @wraps(f)
    def w(*a, **k):
        try:
            g.uid = signer.loads(request.headers.get("Authorization", "").replace("Bearer ", ""), max_age=86400 * 7)
        except BadSignature:
            return jsonify(error="Unauthorized"), 401
        return f(*a, **k)
    return w

_last_call = [0.0]
_call_lock = threading.Lock()

def _gemini_once(model, prompt, tokens, as_json):
    gen = {"maxOutputTokens": tokens + 2000}
    if as_json:
        gen["responseMimeType"] = "application/json"
    if "2.5-flash" in model:
        gen["thinkingConfig"] = {"thinkingBudget": 0}   # keeps the whole answer budget for the text
    body = {"contents": [{"parts": [{"text": prompt}]}], "generationConfig": gen}
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    r = None
    for attempt in range(3):
        try:
            r = requests.post(url, headers={"x-goog-api-key": GEMINI_KEY}, json=body, timeout=GEMINI_TIMEOUT)
        except requests.exceptions.Timeout:
            if attempt < 1:
                continue
            return None, 0, f"Google did not answer within {GEMINI_TIMEOUT} seconds (slow reply or weak internet). Try again, or raise GEMINI_TIMEOUT_SECONDS in .env."
        except requests.exceptions.RequestException as e:
            if attempt < 2:
                time.sleep(3)
                continue
            return None, 0, f"Cannot reach Google ({type(e).__name__}). Check this computer's internet connection (or VPN / DNS) and try again."
        if r.status_code in (429, 500, 503) and attempt < 2:
            time.sleep(8 * (attempt + 1))
            continue
        break
    if not r.ok:
        try:
            msg = r.json()["error"]["message"]
        except Exception:
            msg = r.text[:200]
        return None, r.status_code, msg
    d = r.json()
    cand = (d.get("candidates") or [{}])[0]
    text = "".join(p.get("text", "") for p in (cand.get("content") or {}).get("parts", []))
    if not text.strip():
        return None, 200, f"no text returned (reason: {cand.get('finishReason') or d.get('promptFeedback')})"
    return text, 200, ""

def _gemini(prompt, tokens, as_json):
    if not GEMINI_KEY:
        raise AIError("GEMINI_API_KEY is missing in .env")
    with _call_lock:                                     # one call at a time, spaced out
        wait = GEMINI_GAP - (time.time() - _last_call[0])
        if wait > 0:
            time.sleep(wait)
        err = ""
        for model in [GEMINI_MODEL] + GEMINI_FALLBACKS:
            text, code, msg = _gemini_once(model, prompt, tokens, as_json)
            if text:
                _last_call[0] = time.time()
                return text
            err = f"Gemini {code} ({model}): {msg}"
            print("[AI]", err)
            if code not in (404, 429, 500, 503):         # a key, request or network problem: other models won't help
                break
        _last_call[0] = time.time()
    raise AIError(err)

def _parse_json(t):
    t = re.sub(r"^```(?:json)?|```$", "", t.strip(), flags=re.M).strip()
    try:
        return json.loads(t)
    except ValueError:
        for o, c in (("[", "]"), ("{", "}")):
            i, j = t.find(o), t.rfind(c)
            if i != -1 and j > i:
                try:
                    return json.loads(t[i:j + 1])
                except ValueError:
                    pass
    raise AIError("The AI did not return valid JSON. Please try again.")

def ask(prompt, tokens=4000, as_json=False):
    if PROVIDER == "gemini":
        t = _gemini(prompt, tokens, as_json)
    else:
        r = ai.messages.create(model=MODEL, max_tokens=tokens, messages=[{"role": "user", "content": prompt}])
        t = r.content[0].text
    return _parse_json(t) if as_json else t

LEVEL_RULES = {
    "BSc": "practical project with implementation",
    "MSc": "deeper research with rigorous evaluation",
    "PhD": "original contribution, clear research gap, rigorous methodology",
}
def ctx(d):
    return f"Level: {d['level']} ({LEVEL_RULES.get(d['level'], '')}); Department: {d['department']}; University: {d['university']} (Nigeria)."

def pctx(p):
    # Read the fields directly: after a commit the object's __dict__ is emptied, which caused KeyError: 'level'
    return ctx({"level": p.level, "department": p.department, "university": p.university})

SECTIONS = ["Chapter 1: Introduction", "Chapter 2: Literature Review", "Chapter 3: Methodology",
            "Chapter 4: Data Analysis & Results", "Chapter 5: Conclusion & Recommendation",
            "References", "System Design & Diagrams", "Defense Preparation"]

# ---------- auth ----------
@app.post("/api/auth/register")
def register():
    d = request.json
    if User.query.filter_by(email=d["email"].lower()).first():
        return jsonify(error="Email already registered"), 400
    u = User(name=d["name"], email=d["email"].lower(), pw=generate_password_hash(d["password"]))
    db.session.add(u); db.session.commit()
    return jsonify(token=signer.dumps(u.id), name=u.name)

@app.post("/api/auth/login")
def login():
    d = request.json
    u = User.query.filter_by(email=d["email"].lower()).first()
    if not u or not check_password_hash(u.pw, d["password"]):
        return jsonify(error="Invalid credentials"), 401
    return jsonify(token=signer.dumps(u.id), name=u.name)

# ---------- AI ----------
@app.post("/api/topics")
@auth
def topics():
    d = request.json
    out = ask(f"""{ctx(d)} Interest: {d.get('interest','any')}.
Generate 20 current, supervisor-friendly project topics ranked easiest to hardest. Return ONLY a JSON array of objects:
title, description (2 sentences), difficulty (Easy|Medium|Hard), relevance (High|Medium), tags (array), research_gap, technologies (array), est_time.""", 6000, True)
    return jsonify(topics=out)

@app.post("/api/check-topic")
@auth
def check_topic():
    d = request.json
    out = ask(f"""{ctx(d)} Evaluate this proposed topic: "{d['topic']}".
Return ONLY JSON: verdict (Strong|Needs modification|Too broad/weak), strength, research_gap, feasibility, suggestions (array of 3 better titles).""", 1500, True)
    return jsonify(out)

@app.post("/api/projects")
@auth
def create_project():
    d = request.json
    intro = ask(f"{ctx(d)} Topic: {d['topic']['title']}. Write a Chapter 1 Introduction (background, problem statement, aim, objectives, research questions, scope) in formal academic style. Plain text with numbered headings.", 2500)
    p = Project(user_id=g.uid, level=d["level"], department=d["department"], university=d["university"],
                title=d["topic"]["title"], topic=d["topic"], content={SECTIONS[0]: intro}, paid=not PAY_REQUIRED)
    db.session.add(p); db.session.commit()
    return jsonify(preview(p))

def preview(p):
    # Server-side protection: unpaid users only ever receive a short excerpt
    c = p.content.get(SECTIONS[0], "")
    return {"id": p.id, "title": p.title, "level": p.level, "topic": p.topic, "sections": SECTIONS,
            "excerpt": c[:700] + ("…" if len(c) > 700 else ""), "paid": p.paid,
            "status": p.status, "progress": p.progress, "price": PRICE, "payment_required": PAY_REQUIRED}

@app.get("/api/projects/<int:pid>")
@auth
def get_project(pid):
    p = Project.query.filter_by(id=pid, user_id=g.uid).first_or_404()
    return jsonify(preview(p))

# ---------- payments (Paystack) ----------
@app.post("/api/projects/<int:pid>/pay")
@auth
def pay(pid):
    p = Project.query.filter_by(id=pid, user_id=g.uid).first_or_404()
    if p.paid:
        return jsonify(paid=True)
    u = db.session.get(User, g.uid)
    ref = uuid.uuid4().hex
    db.session.add(Payment(reference=ref, project_id=p.id, amount=PRICE)); db.session.commit()
    if not PSK:
        return jsonify(error="Payments are not set up: PAYSTACK_SECRET is empty in .env."), 502
    try:
        r = requests.post("https://api.paystack.co/transaction/initialize",
            headers={"Authorization": f"Bearer {PSK}"},
            json={"email": u.email, "amount": PRICE * 100, "reference": ref, "currency": "NGN",
                  "callback_url": f"{os.getenv('FRONTEND_URL')}/?project={p.id}"}, timeout=30).json()
    except requests.RequestException as e:
        print(f"[PAYSTACK] network error: {e}")
        return jsonify(error="Could not reach Paystack. Check your internet and try again."), 502
    if not r.get("status") or not r.get("data"):
        msg = r.get("message", "unknown error")
        print(f"[PAYSTACK] initialize failed: {msg}")
        return jsonify(error=f"Paystack said: {msg}"), 502
    return jsonify(url=r["data"]["authorization_url"], reference=ref)

def settle(ref):
    pay_ = Payment.query.filter_by(reference=ref).first()
    if not pay_ or pay_.status == "success":
        return pay_
    r = requests.get(f"https://api.paystack.co/transaction/verify/{ref}", headers={"Authorization": f"Bearer {PSK}"}).json()
    if r.get("data", {}).get("status") == "success" and r["data"]["amount"] >= pay_.amount * 100:
        pay_.status = "success"; db.session.get(Project, pay_.project_id).paid = True; db.session.commit()
    return pay_

@app.get("/api/payments/verify/<ref>")
@auth
def verify(ref):
    p = settle(ref)
    return jsonify(paid=bool(p and p.status == "success"))

@app.post("/api/paystack/webhook")
def webhook():
    sig = hmac.new(PSK.encode(), request.data, hashlib.sha512).hexdigest()
    if not hmac.compare_digest(sig, request.headers.get("x-paystack-signature", "")):
        return "", 401
    e = request.json
    if e["event"] == "charge.success":
        settle(e["data"]["reference"])
    return "", 200

# ---------- diagrams ----------
def png(fig):
    b = io.BytesIO(); fig.savefig(b, format="png", dpi=130, bbox_inches="tight"); plt.close(fig)
    return base64.b64encode(b.getvalue()).decode()

def draw_arch(nodes):
    fig, ax = plt.subplots(figsize=(2.4 * len(nodes), 2)); ax.axis("off"); ax.set_xlim(0, len(nodes)); ax.set_ylim(0, 1)
    for i, n in enumerate(nodes):
        ax.text(i + .5, .5, n.replace(" ", "\n", 1) if len(n) > 12 else n, ha="center", va="center", fontsize=9,
                bbox=dict(boxstyle="round,pad=.6", fc="#fff7ed", ec="#f97316"))
        if i: ax.annotate("", xy=(i + .17, .5), xytext=(i - .17, .5), arrowprops=dict(arrowstyle="->", color="#444"))
    return png(fig)

def draw_er(ents):
    fig, ax = plt.subplots(figsize=(3 * len(ents), 3.6)); ax.axis("off"); ax.set_xlim(0, len(ents)); ax.set_ylim(0, 1)
    for i, e in enumerate(ents):
        t = e["name"] + "\n" + "-" * 16 + "\n" + "\n".join(e["fields"][:7])
        ax.text(i + .5, .5, t, ha="center", va="center", fontsize=9, family="monospace",
                bbox=dict(boxstyle="square,pad=.6", fc="#fff", ec="#111"), zorder=3)
        if i: ax.plot([i - .08, i + .08], [.5, .5], color="#f97316", lw=2, zorder=1)
    return png(fig)

# ---------- background generation job ----------
def run_job(pid):
    with app.app_context():
        p = db.session.get(Project, pid)
        try:
            content = dict(p.content); n = len(SECTIONS) + 1
            for i, s in enumerate(SECTIONS):
                if s not in content:
                    content[s] = ask(f"{pctx(p)} Topic: {p.title}. Write the full '{s}' section in formal academic style. "
                                     "Use only real, verifiable sources; mark any citation the student must verify with [VERIFY]. Plain text.", 3500)
                p.content = dict(content); p.progress = int((i + 1) / n * 90); db.session.commit()
            spec = ask(f"{pctx(p)} Topic: {p.title}. Return ONLY JSON: architecture (3-6 component names in flow order; "
                       "for non-software topics use the research framework stages), entities (2-4 objects {name, fields[]}; "
                       "for non-software topics use key variables).", 1200, True)
            p.diagrams = {"architecture": draw_arch(spec["architecture"]), "er": draw_er(spec["entities"])}
            p.status, p.progress = "done", 100
        except Exception as e:
            import traceback; traceback.print_exc(); print(f"[JOB {pid} FAILED] {e}")
            db.session.rollback(); p = db.session.get(Project, pid); p.status, p.error = "error", str(e)[:200]
        db.session.commit()

def mine(pid):
    return Project.query.filter_by(id=pid, user_id=g.uid).first_or_404()

def paid_only(p):
    return None if p.paid else (jsonify(error="Payment required"), 402)

@app.post("/api/projects/<int:pid>/generate")
@auth
def generate(pid):
    p = mine(pid)
    if (r := paid_only(p)): return r
    if p.status != "running":
        p.status, p.progress, p.error = "running", 0, None; db.session.commit()
        threading.Thread(target=run_job, args=(pid,), daemon=True).start()
    return jsonify(status="running")

@app.get("/api/projects/<int:pid>/content")
@auth
def content(pid):
    p = mine(pid)
    if (r := paid_only(p)): return r
    return jsonify(sections=p.content, diagrams=p.diagrams, status=p.status, progress=p.progress, error=p.error)

@app.put("/api/projects/<int:pid>/sections")
@auth
def save_section(pid):
    p = mine(pid)
    if (r := paid_only(p)): return r
    d = request.json
    if d["section"] not in SECTIONS: return jsonify(error="Unknown section"), 400
    c = dict(p.content); c[d["section"]] = d["text"]; p.content = c; db.session.commit()
    return jsonify(ok=True)

@app.post("/api/projects/<int:pid>/rewrite")
@auth
def rewrite(pid):
    p = mine(pid)
    if (r := paid_only(p)): return r
    d = request.json
    inst = {"rewrite": "Rewrite this more clearly and academically, keeping the meaning.",
            "grammar": "Fix grammar, spelling and punctuation only. Do not change the content.",
            "expand": "Expand this with more depth, keeping the same structure."}[d["action"]]
    return jsonify(text=ask(f"{pctx(p)} Topic: {p.title}. {inst} Return only the revised text.\n\n{d['text']}", 3500))

@app.get("/api/projects/<int:pid>/download")
@auth
def download(pid):
    p = mine(pid)
    if (r := paid_only(p)): return r
    if p.status != "done": return jsonify(error="Project is not ready yet"), 409
    doc = Document(); doc.add_heading(p.title, 0)
    for s in SECTIONS:
        doc.add_heading(s, 1)
        for para in p.content[s].split("\n"):
            if para.strip(): doc.add_paragraph(para.strip())
        if s.startswith("System Design") and p.diagrams:
            for k, cap in (("architecture", "Figure 1: System / Research Framework"), ("er", "Figure 2: Entity / Variable Diagram")):
                doc.add_picture(io.BytesIO(base64.b64decode(p.diagrams[k])), width=None); doc.add_paragraph(cap)
    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{re.sub(r'[^A-Za-z0-9]+','_',p.title)[:60]}.docx")

# ---------- projects list, library, dashboard ----------
@app.get("/api/projects")
@auth
def my_projects():
    ps = Project.query.filter_by(user_id=g.uid).order_by(Project.id.desc()).all()
    return jsonify([{"id": p.id, "title": p.title, "level": p.level, "status": p.status, "progress": p.progress, "paid": p.paid} for p in ps])

@app.get("/api/library")
@auth
def library():
    return jsonify([{"id": s.id, "topic": s.topic, "level": s.level, "department": s.department, "university": s.university}
                    for s in Saved.query.filter_by(user_id=g.uid).order_by(Saved.id.desc())])

@app.post("/api/library")
@auth
def save_topic():
    d = request.json
    db.session.add(Saved(user_id=g.uid, level=d["level"], department=d["department"], university=d["university"], topic=d["topic"]))
    db.session.commit(); return jsonify(ok=True)

@app.delete("/api/library/<int:sid>")
@auth
def del_saved(sid):
    s = Saved.query.filter_by(id=sid, user_id=g.uid).first_or_404(); db.session.delete(s); db.session.commit()
    return jsonify(ok=True)

@app.get("/api/dashboard")
@auth
def dashboard():
    ps = Project.query.filter_by(user_id=g.uid).order_by(Project.id.desc()).all()
    return jsonify(total=len(ps), done=sum(p.status == "done" for p in ps), running=sum(p.status == "running" for p in ps),
                   saved=Saved.query.filter_by(user_id=g.uid).count(),
                   recent=[{"id": p.id, "title": p.title, "status": p.status, "progress": p.progress} for p in ps[:5]])

# Create the tables whenever the app loads (gunicorn on Railway never runs the block below).
with app.app_context():
    db.create_all()

if __name__ == "__main__":
    app.run(debug=True, port=5000, threaded=True)